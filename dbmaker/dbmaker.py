import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

NAVY = RGBColor(0x1F, 0x38, 0x64)
TEAL = RGBColor(0x2E, 0x86, 0xAB)
GRAY = RGBColor(0x55, 0x55, 0x55)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def add_header(doc, title):
    p_title = doc.add_paragraph()
    r_title = p_title.add_run(title)
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = NAVY
    p_title.paragraph_format.space_after = Pt(2)

    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(2)
    r_c = p_meta.add_run('Course: ')
    r_c.bold = True
    p_meta.add_run('CSE4104 — Web Engineering Lab | ')
    r_s = p_meta.add_run('Section: ')
    r_s.bold = True
    p_meta.add_run('7A | ')
    r_t = p_meta.add_run('Team: ')
    r_t.bold = True
    p_meta.add_run('CSE4104-7A-T02')

    p_proj = doc.add_paragraph()
    p_proj.paragraph_format.space_after = Pt(8)
    r_p = p_proj.add_run('Project: ')
    r_p.bold = True
    p_proj.add_run('AI Sales Analytics Dashboard (SalesIQ) | ')
    r_i = p_proj.add_run('Institution: ')
    r_i.bold = True
    p_proj.add_run('Northern University of Business and Technology, Khulna')

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(14)
    r_div = p_div.add_run('―' * 60)
    r_div.font.color.rgb = NAVY
    r_div.font.bold = True

def add_team_table(doc):
    h = doc.add_heading('Team Information', level=2)
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(4)
    table = doc.add_table(rows=5, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ['Role', 'Name', 'Student ID']
    for idx, text in enumerate(headers):
        cell = table.cell(0, idx)
        set_cell_background(cell, '1F3864')
        p = cell.paragraphs[0]
        r = p.add_run(text)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(9.5)
    
    rows_data = [
        ['Team Leader', 'Mohaimeen Islam Pial', '11230121094'],
        ['Frontend Developer', 'Sk Mesbaul Arefin', '11230121077'],
        ['Backend Developer', 'Sumaiya Akter', '11230121081'],
        ['Database Manager', 'Afia Maliha Priota', '11230121090'],
    ]
    for r_idx, row in enumerate(rows_data, 1):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            if r_idx % 2 == 1:
                set_cell_background(cell, 'F1F5F9')
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9)
            if c_idx == 0:
                r.font.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

def add_screenshot_box(doc, title, hint):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    r_t = p.add_run(f'📷 {title}')
    r_t.font.bold = True
    r_t.font.color.rgb = TEAL
    r_t.font.size = Pt(11)

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, 'F8FAFC')
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''<w:tcBorders {nsdecls("w")}>
        <w:top w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>
        <w:left w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>
        <w:bottom w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>
        <w:right w:val="dashed" w:sz="8" w:space="0" w:color="94A3B8"/>
    </w:tcBorders>''')
    tcPr.append(tcBorders)

    cp = cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(36)
    cp.paragraph_format.space_after = Pt(6)
    r_h = cp.add_run('[ ATTACH SCREENSHOT HERE ]')
    r_h.font.bold = True
    r_h.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    r_h.font.size = Pt(12)

    cp2 = cell.add_paragraph()
    cp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp2.paragraph_format.space_after = Pt(36)
    r_desc = cp2.add_run(hint)
    r_desc.font.italic = True
    r_desc.font.size = Pt(9.5)
    r_desc.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

# -------------------------------------------------------------
# 1. CSE4104-7A-T02_TestCases.docx
# -------------------------------------------------------------
doc1 = docx.Document()
add_header(doc1, 'Test Cases Specification Document')
add_team_table(doc1)

doc1.add_heading('1. Test Case Execution Matrix', level=2)
tc_table = doc1.add_table(rows=16, cols=5)
tc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ['Test ID', 'Module', 'Test Scenario & Input', 'Expected Result', 'Status']
for idx, text in enumerate(headers):
    cell = tc_table.cell(0, idx)
    set_cell_background(cell, '1F3864')
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    r.font.size = Pt(8.5)

tc_rows = [
    ['TC-AUTH-01', 'Auth', 'Valid user login (admin@test.com / admin123)', 'HTTP 200 OK, JWT Bearer token & user payload returned', 'Passed'],
    ['TC-AUTH-02', 'Auth', 'Login with invalid password', 'HTTP 401 Unauthorized with error feedback', 'Passed'],
    ['TC-AUTH-04', 'Auth', 'Login with empty fields', 'HTTP 422 Unprocessable Content validation error', 'Passed'],
    ['TC-AUTH-07', 'Auth', 'Admin registration without invite code', 'HTTP 400 Bad Request, blocks privilege escalation', 'Passed'],
    ['TC-AUTH-09', 'RBAC', 'Viewer requesting admin endpoint /api/users', 'HTTP 403 Forbidden with role restriction message', 'Passed'],
    ['TC-SALE-01', 'Sales', 'Paginated sales list (page=1, size=2)', 'HTTP 200 OK with exact total and page count metadata', 'Passed'],
    ['TC-SALE-02', 'Sales', 'Filter sales by category (Furniture)', 'HTTP 200 OK returning only Furniture records', 'Passed'],
    ['TC-SALE-03', 'Sales', 'Free-text search (Laptop)', 'HTTP 200 OK matching product/region names', 'Passed'],
    ['TC-SALE-04', 'Sales', 'Manager creates sale record', 'HTTP 201 Created, auto-computed total_price', 'Passed'],
    ['TC-SALE-05', 'RBAC', 'Viewer attempts to create sale', 'HTTP 403 Forbidden, view-only enforcement', 'Passed'],
    ['TC-SALE-08', 'Sales', 'Admin deletes sale record', 'HTTP 200 OK, record safely deleted from DB', 'Passed'],
    ['TC-ANL-01', 'Analytics', 'Calculate Summary KPIs', 'HTTP 200 OK, total revenue, orders, AOV, growth %', 'Passed'],
    ['TC-AI-01', 'AI Engine', 'Generate AI business insights', 'HTTP 200 OK, executive summary & recommendations', 'Passed'],
    ['TC-AI-02', 'AI ML', 'Revenue forecast for next month', 'HTTP 200 OK, linear regression prediction & confidence bounds', 'Passed'],
    ['TC-AI-07', 'AI Fallback', 'Gemini 503 / Network spike resilience', 'Automatic fallback to GroundedAIEngine with zero downtime', 'Passed'],
]
for r_idx, row in enumerate(tc_rows, 1):
    for c_idx, val in enumerate(row):
        cell = tc_table.cell(r_idx, c_idx)
        if r_idx % 2 == 1:
            set_cell_background(cell, 'F1F5F9')
        p = cell.paragraphs[0]
        r = p.add_run(val)
        r.font.size = Pt(8)
        if c_idx == 0:
            r.font.bold = True
        if c_idx == 4:
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x16, 0xA3, 0x4A)

doc1.add_paragraph().paragraph_format.space_after = Pt(10)
doc1.add_heading('2. Evidence of Test Case Execution', level=2)
add_screenshot_box(doc1, 'Screenshot 1: Automated Pytest Suite Execution', 'Capture of terminal output running python -m pytest -v showing 24 passed test cases in green.')
add_screenshot_box(doc1, 'Screenshot 2: API Endpoints Verification (Postman / Swagger /docs)', 'Capture of successful API requests (POST /api/auth/login and GET /api/sales) with HTTP 200 OK.')
add_screenshot_box(doc1, 'Screenshot 3: RBAC Authorization & Security Guard', 'Capture showing HTTP 403 Forbidden when Viewer attempts an unauthorized Admin/Manager action.')
doc1.save('CSE4104-7A-T02_TestCases.docx')
print('Generated CSE4104-7A-T02_TestCases.docx')

# -------------------------------------------------------------
# 2. CSE4104-7A-T02_BugReport.docx
# -------------------------------------------------------------
doc2 = docx.Document()
add_header(doc2, 'Bug Tracking & Defect Resolution Report')
add_team_table(doc2)

doc2.add_heading('1. Defect Tracking Matrix', level=2)
bug_table = doc2.add_table(rows=8, cols=6)
bug_table.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ['Bug ID', 'Module', 'Description', 'Severity', 'Status', 'Resolution']
for idx, text in enumerate(headers):
    cell = bug_table.cell(0, idx)
    set_cell_background(cell, '1F3864')
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    r.font.size = Pt(8.5)

bug_rows = [
    ['BUG-01', 'Auth', 'Empty login fields submitted without client/server validation', 'Medium', 'Fixed', 'Added Pydantic schema validation & client-side helper text alerts.'],
    ['BUG-02', 'RBAC', 'Viewer role was able to post sale records', 'High', 'Fixed', 'Enforced require_manager_or_admin dependency on POST /api/sales.'],
    ['BUG-03', 'RBAC', 'Manager role could permanently delete sales records', 'High', 'Fixed', 'Restricted DELETE /api/sales/{id} strictly to require_admin role.'],
    ['BUG-04', 'AI Layer', 'External Gemini API 503 high-demand spike caused crash', 'High', 'Fixed', 'Integrated seamless offline fallback to GroundedAIEngine (zero downtime).'],
    ['BUG-05', 'AI Chat', 'Adversarial SQL injection strings in chatbot input', 'Critical', 'Fixed', 'Implemented deterministic regex router + parameterized ORM queries.'],
    ['BUG-06', 'UI/UX', 'Mobile sidebar drawer did not auto-collapse on navigation', 'Low', 'Fixed', 'Added route change event listener in Sidebar to close drawer on < 768px.'],
    ['BUG-07', 'Sales UI', 'Fast keystrokes in search bar sent repeated redundant requests', 'Low', 'Fixed', 'Added 300ms debounce timer to search input handler in React.'],
]
for r_idx, row in enumerate(bug_rows, 1):
    for c_idx, val in enumerate(row):
        cell = bug_table.cell(r_idx, c_idx)
        if r_idx % 2 == 1:
            set_cell_background(cell, 'F1F5F9')
        p = cell.paragraphs[0]
        r = p.add_run(val)
        r.font.size = Pt(8)
        if c_idx == 0:
            r.font.bold = True
        if c_idx == 4:
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x16, 0xA3, 0x4A)

doc2.add_paragraph().paragraph_format.space_after = Pt(10)
doc2.add_heading('2. Evidence of Bug Identification & Resolution', level=2)
add_screenshot_box(doc2, 'Screenshot 1: Bug Identification / Error Scenario', 'Capture of defect behavior before fix (e.g. 422/503 error handling or form validation state).')
add_screenshot_box(doc2, 'Screenshot 2: Bug Resolution & Successful Retesting', 'Capture of corrected behavior after fix (e.g. graceful validation message or AI fallback response).')
add_screenshot_box(doc2, 'Screenshot 3: GitHub Commit History for Bug Fixes', 'Capture of GitHub commit log showing descriptive commit messages.')
doc2.save('CSE4104-7A-T02_BugReport.docx')
print('Generated CSE4104-7A-T02_BugReport.docx')

# -------------------------------------------------------------
# 3. CSE4104-7A-T02_TestingReport.docx
# -------------------------------------------------------------
doc3 = docx.Document()
add_header(doc3, 'Testing and Quality Assurance Report')
add_team_table(doc3)

doc3.add_heading('1. Executive Summary & Testing Approach', level=2)
p_appr = doc3.add_paragraph(
    'The AI Sales Analytics Dashboard (SalesIQ) underwent comprehensive full-stack testing across all eight '
    'mandatory assessment areas: Functional Testing, Authentication & RBAC, API Testing, Database Integrity, '
    'AI Quality & Resilience Evaluation, UI/UX, Responsive Design, and Security Analysis. Testing was conducted '
    'using automated test runners (Pytest, HTTPX), browser testing across viewports, and static code security audits.'
)
p_appr.paragraph_format.space_after = Pt(6)

doc3.add_heading('2. Automated Test Suite & Regression Summary', level=2)
p_reg = doc3.add_paragraph(
    '• Automated Pytest Suite: 24 unit/integration tests executed in 34.61s with 100% pass rate (24 passed, 0 failed).\n'
    '• Frontend Production Build: Clean Vite compilation with 0 errors across 2,406 transformed modules.\n'
    '• Database Integrity: Verified constraints, foreign keys (ON DELETE RESTRICT/SET NULL), and indexing.'
)
p_reg.paragraph_format.space_after = Pt(6)
add_screenshot_box(doc3, 'Screenshot 1: Automated Test Suite & Regression Run', 'Terminal screenshot of python -m pytest with 24 passed tests.')

doc3.add_heading('3. AI Quality Evaluation (10-Input Benchmark)', level=2)
ai_table = doc3.add_table(rows=11, cols=4)
ai_table.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ['#', 'Input Category', 'Test Query', 'Result & Assessment']
for idx, text in enumerate(headers):
    cell = ai_table.cell(0, idx)
    set_cell_background(cell, '1F3864')
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    r.font.size = Pt(8.5)

ai_rows = [
    ['1', 'Normal Query 1', 'What was our best performing product?', 'PASS: Correctly identified top revenue product with real figures.'],
    ['2', 'Normal Query 2', 'Show me sales performance for Dhaka region', 'PASS: Accurately filtered regional revenue and order count.'],
    ['3', 'Very Short Input', 'sales?', 'PASS: Handled gracefully; rendered overall revenue summary.'],
    ['4', 'Very Long Input', 'Detailed multi-KPI query across products, regions, anomalies', 'PASS: Processed full context; delivered structured executive summary.'],
    ['5', 'Irrelevant Input', 'What is the recipe for chocolate cake?', 'PASS: Politely redirected user to sales and dashboard topics.'],
    ['6', 'Ambiguous Input', 'How are things?', 'PASS: Conversational persona response summarizing system health.'],
    ['7', 'Adversarial / SQLi', '\'; DROP TABLE sales; -- SELECT * FROM users', 'PASS: Safely sanitized; no SQL execution occurred.'],
    ['8', 'Repeated Input', 'total revenue total revenue total revenue', 'PASS: Deduplicated query; returned total revenue figure.'],
    ['9', 'Empty / Whitespace', '[Whitespace only]', 'PASS: Safely handled; prompted user with helpful guidance.'],
    ['10', 'Definitional', 'What does AOV mean and how is it calculated?', 'PASS: Returned precise business formula (Total Revenue / Total Orders).'],
]
for r_idx, row in enumerate(ai_rows, 1):
    for c_idx, val in enumerate(row):
        cell = ai_table.cell(r_idx, c_idx)
        if r_idx % 2 == 1:
            set_cell_background(cell, 'F1F5F9')
        p = cell.paragraphs[0]
        r = p.add_run(val)
        r.font.size = Pt(8)
        if c_idx == 0:
            r.font.bold = True
        if c_idx == 3:
            r.font.bold = True

doc3.add_paragraph().paragraph_format.space_after = Pt(8)
add_screenshot_box(doc3, 'Screenshot 2: AI Insights & Natural Language Chatbot in Action', 'Capture of the AI Sales Chatbot answering sales queries and generating period insights.')

doc3.add_heading('4. Responsive Design & Security Verification', level=2)
p_sec = doc3.add_paragraph(
    '• Responsive Testing: Verified on Desktop (1280px+), Tablet (768px), and Mobile (375px) viewports with dynamic charts and mobile drawer navigation.\n'
    '• Security Checks: Passwords hashed with bcrypt (12 rounds), JWT token auth, zero secrets exposed in Git, 100% parameterized ORM queries.'
)
p_sec.paragraph_format.space_after = Pt(6)
add_screenshot_box(doc3, 'Screenshot 3: Responsive Layout Testing (Mobile 375px & Desktop 1280px)', 'Side-by-side or stacked capture of dashboard in Chrome DevTools Device Mode.')
add_screenshot_box(doc3, 'Screenshot 4: Database & Security Verification', 'Capture of Supabase/SQLite database showing hashed passwords ($2b$12$) and foreign key relations.')

doc3.save('CSE4104-7A-T02_TestingReport.docx')
print('Generated CSE4104-7A-T02_TestingReport.docx')